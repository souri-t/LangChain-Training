"""
LangGraphを利用したPowerPoint解析AIエージェント

機能: PowerPointファイルからテキストを抽出し、日本語→英語に翻訳して上書き保存
フロー: PPTXファイル読込 → スライド単位でテキスト抽出・翻訳 → 上書き保存

例:
入力: sample.pptx (日本語テキスト含む)
→ 各スライドのテキストを抽出
→ LLMで日→英に翻訳
→ 英文で上書きして保存
"""

import os
import operator
from typing import TypedDict, Annotated, List, Dict, Any
from dotenv import load_dotenv
from langchain_core.messages import BaseMessage, HumanMessage, SystemMessage
from langchain_openai import ChatOpenAI
from langgraph.graph import StateGraph, END
from langgraph.prebuilt import ToolNode
from pptx import Presentation

# PowerPoint操作ツールのインポート
from tools.powerpoint_tools import tools

# .envファイルから環境変数を読み込み
load_dotenv()

# 環境変数の取得
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
OPENAI_API_BASE = os.getenv("OPENAI_API_BASE", "https://api.openai.com/v1")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")


# =============================================================================
# 1. 状態定義
# =============================================================================

class PowerPointTranslationState(TypedDict):
    """PowerPoint翻訳エージェントの状態を管理する"""
    file_path: str  # PowerPointファイルのパス
    slide_count: int  # 総スライド数
    current_slide_index: int  # 現在処理中のスライドインデックス
    translated_count: int  # 翻訳完了したスライド数
    current_texts: List[Dict[str, Any]]  # 現在のスライドのテキスト
    current_translations: List[Dict[str, str]]  # 現在のスライドの翻訳結果
    messages: Annotated[List[BaseMessage], operator.add]  # LLMとのメッセージ履歴
    errors: List[str]  # エラーメッセージのリスト


# =============================================================================
# 2. ノード定義
# =============================================================================

def load_ppt_node(state: PowerPointTranslationState) -> Dict[str, Any]:
    """
    PowerPointファイルを読み込み、スライド数を取得する
    """
    file_path = state["file_path"]
    
    print(f"\n{'='*60}")
    print(f"PowerPointファイルを読み込んでいます: {file_path}")
    print(f"{'='*60}\n")
    
    try:
        prs = Presentation(file_path)
        slide_count = len(prs.slides)
        
        print(f"✓ 読み込み成功: {slide_count}スライド")
        
        return {
            "slide_count": slide_count,
            "current_slide_index": 0,
            "translated_count": 0,
            "errors": []
        }
    except Exception as e:
        error_msg = f"ファイル読み込みエラー: {str(e)}"
        print(f"✗ {error_msg}")
        return {
            "slide_count": 0,
            "errors": [error_msg]
        }


def extract_and_translate_node(state: PowerPointTranslationState) -> Dict[str, Any]:
    """
    現在のスライドからテキストを抽出し、LLMで翻訳する
    """
    file_path = state["file_path"]
    slide_index = state["current_slide_index"]
    
    print(f"\n--- スライド {slide_index + 1} の処理 ---")
    
    try:
        # テキスト抽出
        prs = Presentation(file_path)
        slide = prs.slides[slide_index]
        
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append({
                    "text": shape.text,
                    "shape_id": shape.shape_id
                })
        
        print(f"テキスト抽出: {len(texts)}個のテキストボックス")
        
        if not texts:
            print("→ テキストなし、スキップします")
            return {
                "current_texts": [],
                "current_translations": [],
                "current_slide_index": slide_index + 1,
                "translated_count": state["translated_count"] + 1
            }
        
        # LLMで翻訳
        llm = ChatOpenAI(
            model=OPENAI_MODEL,
            temperature=0,
            openai_api_key=OPENAI_API_KEY,
            openai_api_base=OPENAI_API_BASE,
            max_retries=3,
            timeout=60
        )
        
        translations = []
        for item in texts:
            original_text = item["text"]
            
            # 翻訳プロンプト
            translation_prompt = f"""以下の日本語テキストを自然な英語に翻訳してください。
翻訳結果のみを出力し、説明や追加のコメントは不要です。

日本語テキスト:
{original_text}

英語翻訳:"""
            
            try:
                response = llm.invoke([HumanMessage(content=translation_prompt)])
                translated_text = response.content.strip()
                
                translations.append({
                    "shape_id": item["shape_id"],
                    "original": original_text,
                    "translated": translated_text
                })
                
                print(f"  原文: {original_text[:50]}...")
                print(f"  翻訳: {translated_text[:50]}...")
                
            except Exception as e:
                error_msg = f"翻訳エラー (shape_id: {item['shape_id']}): {str(e)}"
                print(f"  ✗ {error_msg}")
                # エラーの場合は元のテキストを保持
                translations.append({
                    "shape_id": item["shape_id"],
                    "original": original_text,
                    "translated": original_text  # 翻訳失敗時は元のテキストを使用
                })
        
        return {
            "current_texts": texts,
            "current_translations": translations,
            "messages": []
        }
        
    except Exception as e:
        error_msg = f"スライド {slide_index + 1} の処理エラー: {str(e)}"
        print(f"✗ {error_msg}")
        errors = state.get("errors", [])
        errors.append(error_msg)
        return {
            "current_texts": [],
            "current_translations": [],
            "errors": errors
        }


def save_translations_node(state: PowerPointTranslationState) -> Dict[str, Any]:
    """
    翻訳結果をPowerPointファイルに書き込む
    """
    file_path = state["file_path"]
    slide_index = state["current_slide_index"]
    translations = state["current_translations"]
    
    if not translations:
        # 翻訳がない場合はスキップ
        return {
            "current_slide_index": slide_index + 1,
            "translated_count": state["translated_count"] + 1
        }
    
    try:
        prs = Presentation(file_path)
        slide = prs.slides[slide_index]
        
        updated_count = 0
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                for trans in translations:
                    if trans.get("shape_id") == shape.shape_id:
                        # テキストフレームの全ての段落をクリアして新しいテキストを設定
                        text_frame = shape.text_frame
                        text_frame.clear()
                        text_frame.text = trans.get("translated", "")
                        updated_count += 1
        
        # ファイルを保存
        prs.save(file_path)
        
        print(f"✓ スライド {slide_index + 1} を更新しました ({updated_count}個のテキスト)")
        
        return {
            "current_slide_index": slide_index + 1,
            "translated_count": state["translated_count"] + 1
        }
        
    except Exception as e:
        error_msg = f"スライド {slide_index + 1} の保存エラー: {str(e)}"
        print(f"✗ {error_msg}")
        errors = state.get("errors", [])
        errors.append(error_msg)
        return {
            "current_slide_index": slide_index + 1,
            "translated_count": state["translated_count"] + 1,
            "errors": errors
        }


def should_continue_translation(state: PowerPointTranslationState) -> str:
    """
    翻訳を続けるか、終了するかを判断
    """
    current_index = state["current_slide_index"]
    slide_count = state["slide_count"]
    
    if current_index < slide_count:
        return "continue"
    else:
        return "end"


# =============================================================================
# 3. グラフ構築
# =============================================================================

def create_powerpoint_translation_graph():
    """PowerPoint翻訳エージェントのグラフを作成"""
    
    # グラフの初期化
    workflow = StateGraph(PowerPointTranslationState)
    
    # ノードの追加
    workflow.add_node("load_ppt", load_ppt_node)
    workflow.add_node("extract_and_translate", extract_and_translate_node)
    workflow.add_node("save_translations", save_translations_node)
    
    # エントリーポイントの設定
    workflow.set_entry_point("load_ppt")
    
    # エッジの追加
    workflow.add_edge("load_ppt", "extract_and_translate")
    workflow.add_edge("extract_and_translate", "save_translations")
    
    # 条件付きエッジ: 全スライド処理完了まで繰り返し
    workflow.add_conditional_edges(
        "save_translations",
        should_continue_translation,
        {
            "continue": "extract_and_translate",
            "end": END
        }
    )
    
    # グラフのコンパイル
    app = workflow.compile()
    
    return app


# =============================================================================
# 4. エージェント実行
# =============================================================================

def run_powerpoint_translation_agent(file_path: str):
    """
    PowerPoint翻訳エージェントを実行
    
    Args:
        file_path: PowerPointファイルのパス
    
    Returns:
        実行結果の辞書
    """
    # グラフの作成
    app = create_powerpoint_translation_graph()
    
    # 初期状態
    initial_state = {
        "file_path": file_path,
        "slide_count": 0,
        "current_slide_index": 0,
        "translated_count": 0,
        "current_texts": [],
        "current_translations": [],
        "messages": [],
        "errors": []
    }
    
    print(f"\n{'='*60}")
    print(f"PowerPoint翻訳エージェント開始")
    print(f"対象ファイル: {file_path}")
    print(f"{'='*60}\n")
    
    # エージェントの実行
    final_state = None
    for step, state in enumerate(app.stream(initial_state), 1):
        final_state = state
    
    # 最終状態から結果を取得
    result_state = list(final_state.values())[0]
    
    print(f"\n{'='*60}")
    print(f"翻訳完了")
    print(f"処理スライド数: {result_state.get('translated_count', 0)} / {result_state.get('slide_count', 0)}")
    
    if result_state.get("errors"):
        print(f"\nエラー:")
        for error in result_state["errors"]:
            print(f"  - {error}")
    
    print(f"{'='*60}\n")
    
    return {
        "success": len(result_state.get("errors", [])) == 0,
        "translated_count": result_state.get("translated_count", 0),
        "slide_count": result_state.get("slide_count", 0),
        "errors": result_state.get("errors", [])
    }


# =============================================================================
# 5. メイン実行
# =============================================================================

if __name__ == "__main__":
    import sys
    
    # コマンドライン引数からファイルパスを取得
    if len(sys.argv) > 1:
        ppt_file = sys.argv[1]
    else:
        # デフォルトのファイルパスを設定（テスト用）
        ppt_file = "sample.pptx"
        print(f"使用方法: python pp_main.py <PowerPointファイルのパス>")
        print(f"デフォルトファイルを使用します: {ppt_file}\n")
    
    # ファイルの存在確認
    if not os.path.exists(ppt_file):
        print(f"エラー: ファイルが見つかりません: {ppt_file}")
        sys.exit(1)
    
    # エージェントの実行
    result = run_powerpoint_translation_agent(ppt_file)

